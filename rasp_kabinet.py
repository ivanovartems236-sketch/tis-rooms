"""Расписание кабинета ТИС (филиал ДГТУ) на выбранный день.
По кнопке скачивает актуальные файлы со https://stis.su/schedule и показывает,
что идёт в выбранной аудитории (СЗ = спортзал).
Зависимость: pip install openpyxl
"""
import datetime as dt
import io
import os
import re
import ssl
import sys
import threading
import urllib.parse
import urllib.request
import warnings

warnings.filterwarnings("ignore")
import openpyxl

SITE = "https://stis.su/schedule"
SKIP = ("rasp_exam", "zaschita")  # экзамены и ВКР — другой формат
DAYS = ["ПОНЕДЕЛЬНИК", "ВТОРНИК", "СРЕДА", "ЧЕТВЕРГ", "ПЯТНИЦА", "СУББОТА", "ВОСКРЕСЕНЬЕ"]
TYPES = {"л": "Лекция", "пр": "Практика", "лаб": "Лаб."}
ROOM_RE = re.compile(r"^[А-ЯЁA-Z]{1,3}(-\d+[А-Яа-яA-Za-z]?)?$")
SHEET_RE = re.compile(r"(\d{1,2})\.(\d{1,2})\s*-\s*(\d{1,2})\.(\d{1,2})")
CACHE = os.path.join(os.path.dirname(os.path.abspath(sys.argv[0])), "rasp_cache")


# ---------- загрузка ----------
def _get(url):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    try:
        return urllib.request.urlopen(req, timeout=30).read()
    except urllib.error.URLError as e:
        if isinstance(getattr(e, "reason", None), ssl.SSLError):
            ctx = ssl._create_unverified_context()
            return urllib.request.urlopen(req, timeout=30, context=ctx).read()
        raise


def download_all(log=print, folder=CACHE):
    html = _get(SITE).decode("utf-8", "ignore")
    links = []
    for href in re.findall(r'href="([^"]+\.xlsx?)"', html, re.I):
        url = urllib.parse.urljoin(SITE, href)
        if url not in links and not any(s in url for s in SKIP):
            links.append(url)
    os.makedirs(folder, exist_ok=True)
    for f in os.listdir(folder):
        os.remove(os.path.join(folder, f))
    for i, url in enumerate(links):
        name = urllib.parse.unquote(url.split("/")[-1])
        name = f"{i:02d}_" + re.sub(r'[\\/:*?"<>|]', "_", name)
        try:
            with open(os.path.join(folder, name), "wb") as fh:
                fh.write(_get(url))
            log(f"Скачано: {name}")
        except Exception as e:
            log(f"Ошибка {name}: {e}")
    return len(links)


# ---------- разбор ----------
def sheet_contains(title, day):
    m = SHEET_RE.search(title)
    if not m:
        return False
    d1, m1, d2, m2 = map(int, m.groups())
    for y in (day.year - 1, day.year):
        try:
            start = dt.date(y, m1, d1)
            end = dt.date(y if m2 >= m1 else y + 1, m2, d2)
        except ValueError:
            continue
        if start <= day <= end:
            return True
    return False


def parse_cell(v):
    lines = [l.strip() for l in str(v).split("\n") if l.strip()]
    first = lines[0]
    m = re.match(r"^(лаб|пр|л)\.\s*(.*)$", first)
    kind, subj = (TYPES[m.group(1)], m.group(2)) if m else ("", first)
    tail = lines[-1] if len(lines) > 1 else ""
    toks = tail.split()
    room = ""
    if toks and ROOM_RE.match(toks[-1]):
        room = toks.pop()
    teacher = " ".join(toks)
    if len(lines) > 2:
        subj += " " + " ".join(lines[1:-1])
    return kind, subj.strip(), teacher, room


def parse_sheet(ws, src):
    """Возвращает список занятий: dict(day, pair, time, group, kind, subj, teacher, room)."""
    hdr = next((r for r in range(1, 25) if str(ws.cell(r, 2).value or "").strip() == "Пара"), None)
    if not hdr:
        return []
    groups = {c: str(ws.cell(hdr, c).value).strip()
              for c in range(4, ws.max_column + 1) if ws.cell(hdr, c).value}
    merged = {}
    for mr in ws.merged_cells.ranges:
        if mr.min_row > hdr and mr.max_col > mr.min_col:
            merged[(mr.min_row, mr.min_col)] = range(mr.min_col, mr.max_col + 1)
    out, day = [], None
    for r in range(hdr + 1, ws.max_row + 1):
        a = str(ws.cell(r, 1).value or "").strip().upper()
        if a in DAYS:
            day = DAYS.index(a)
        pair = ws.cell(r, 2).value
        try:
            pair = int(str(pair).strip())
        except (TypeError, ValueError):
            continue
        time = str(ws.cell(r, 3).value or "").strip()
        for c in groups:
            v = ws.cell(r, c).value
            if v is None or not str(v).strip():
                continue
            kind, subj, teacher, room = parse_cell(v)
            for cc in merged.get((r, c), [c]):
                if cc in groups:
                    out.append(dict(day=day, pair=pair, time=time, group=groups[cc], kind=kind,
                                    subj=subj, teacher=teacher, room=room, src=src))
    return out


def load_book(path):
    with open(path, "rb") as fh:
        data = fh.read()
    return openpyxl.load_workbook(io.BytesIO(data))  # файлы .xls на сайте на самом деле xlsx


def collect(day, folder=CACHE):
    """Занятия на дату day + множество всех аудиторий."""
    lessons, rooms = [], set()
    if not os.path.isdir(folder):
        return lessons, rooms
    for f in sorted(os.listdir(folder)):
        try:
            wb = load_book(os.path.join(folder, f))
        except Exception:
            continue
        for ws in wb.worksheets:
            if not SHEET_RE.search(ws.title):
                continue
            items = parse_sheet(ws, f)
            rooms.update(i["room"] for i in items if i["room"])
            if sheet_contains(ws.title, day):
                lessons += [i for i in items if i["day"] == day.weekday()]
    return lessons, rooms


def for_room(lessons, room):
    """Группирует по паре: одинаковые занятия разных групп сливаются."""
    res = {}
    for l in lessons:
        if l["room"] != room:
            continue
        key = (l["pair"], l["subj"], l["teacher"], l["kind"])
        if key not in res:
            res[key] = dict(l, groups=[])
        if l["group"] not in res[key]["groups"]:
            res[key]["groups"].append(l["group"])
    return sorted(res.values(), key=lambda x: x["pair"])


def room_label(r):
    return "СЗ (спортзал)" if r == "СЗ" else r


def room_sort(r):
    m = re.match(r"([^\d-]*)-?(\d*)", r)
    return (m.group(1), int(m.group(2) or 0))


# ---------- окно ----------
def run_gui():
    import tkinter as tk
    from tkinter import ttk

    root = tk.Tk()
    root.title("Расписание кабинета — ТИС")
    root.geometry("760x560")
    state = {"day": dt.date.today(), "rooms": []}

    top = ttk.Frame(root, padding=8)
    top.pack(fill="x")
    ttk.Label(top, text="Аудитория:").pack(side="left")
    room_var = tk.StringVar()
    cb = ttk.Combobox(top, textvariable=room_var, width=16, state="readonly")
    cb.pack(side="left", padx=4)
    ttk.Button(top, text="◀", width=3, command=lambda: shift(-1)).pack(side="left", padx=(12, 0))
    day_lbl = ttk.Label(top, width=24, anchor="center")
    day_lbl.pack(side="left")
    ttk.Button(top, text="▶", width=3, command=lambda: shift(1)).pack(side="left")
    ttk.Button(top, text="Сегодня", command=lambda: shift(0)).pack(side="left", padx=4)
    upd = ttk.Button(top, text="⟳ Обновить с сайта")
    upd.pack(side="right")

    txt = tk.Text(root, wrap="word", font=("Segoe UI", 11), padx=10, pady=8)
    txt.pack(fill="both", expand=True)
    txt.tag_configure("h", font=("Segoe UI", 13, "bold"))
    txt.tag_configure("t", font=("Segoe UI", 11, "bold"))
    txt.tag_configure("free", foreground="#888")
    status = ttk.Label(root, text="", padding=4)
    status.pack(fill="x")

    WD = ["пн", "вт", "ср", "чт", "пт", "сб", "вс"]

    def shift(n):
        state["day"] = dt.date.today() if n == 0 else state["day"] + dt.timedelta(days=n)
        render()

    def render(*_):
        d = state["day"]
        day_lbl.config(text=f"{d:%d.%m.%Y} ({WD[d.weekday()]})")
        lessons, rooms = collect(d)
        if rooms and sorted(rooms, key=room_sort) != state["rooms"]:
            state["rooms"] = sorted(rooms, key=room_sort)
            cb["values"] = [room_label(r) for r in state["rooms"]]
        txt.config(state="normal")
        txt.delete("1.0", "end")
        if not state["rooms"]:
            txt.insert("end", "Нет данных. Нажми «Обновить с сайта».")
        elif not room_var.get():
            txt.insert("end", "Выбери аудиторию.")
        else:
            room = room_var.get().split(" ")[0]
            items = for_room(lessons, room)
            txt.insert("end", f"{room_label(room)} — {d:%d.%m.%Y}\n", "h")
            if not lessons:
                txt.insert("end", "\nНа эту дату в расписании нет данных (выходной или неделя ещё не выложена).\n", "free")
            busy = {i["pair"] for i in items}
            times = {}
            for l in lessons:
                times.setdefault(l["pair"], l["time"])
            for p in sorted(set(times) | busy):
                cur = [i for i in items if i["pair"] == p]
                if not cur:
                    txt.insert("end", f"\n{p} пара  {times.get(p, '')}  — свободно\n", "free")
                for i in cur:
                    txt.insert("end", f"\n{p} пара  {i['time']}\n", "t")
                    txt.insert("end", f"{i['kind'] + ': ' if i['kind'] else ''}{i['subj']}\n")
                    txt.insert("end", f"{i['teacher']}\nГруппы: {', '.join(i['groups'])}\n")
            if lessons and not items:
                txt.insert("end", "\nВ этот день занятий в аудитории нет.\n")
        txt.config(state="disabled")

    def update_site():
        upd.config(state="disabled")
        status.config(text="Загрузка с сайта…")

        def work():
            try:
                n = download_all(lambda s: root.after(0, status.config, {"text": s}))
                msg = f"Готово: файлов {n}, {dt.datetime.now():%H:%M}"
            except Exception as e:
                msg = f"Ошибка загрузки: {e}"
            root.after(0, lambda: (status.config(text=msg), upd.config(state="normal"), render()))

        threading.Thread(target=work, daemon=True).start()

    upd.config(command=update_site)
    cb.bind("<<ComboboxSelected>>", render)
    render()
    if not state["rooms"]:
        update_site()
    root.mainloop()


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--cli":  # python rasp_kabinet.py --cli СЗ [папка] [дд.мм.гггг]
        room = sys.argv[2]
        folder = sys.argv[3] if len(sys.argv) > 3 else CACHE
        day = dt.datetime.strptime(sys.argv[4], "%d.%m.%Y").date() if len(sys.argv) > 4 else dt.date.today()
        lessons, rooms = collect(day, folder)
        print("Аудитории:", ", ".join(sorted(rooms, key=room_sort)))
        for i in for_room(lessons, room):
            print(f"{i['pair']} | {i['time']} | {i['kind']} {i['subj']} | {i['teacher']} | {', '.join(i['groups'])}")
    else:
        run_gui()
